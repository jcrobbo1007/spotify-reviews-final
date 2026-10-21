"""Golden-50 evaluation: Jack's labels vs saved pipeline predictions. Offline, no model calls.

50 cases are diagnostic, not a population estimate. Missing/quarantined predictions stay in
agreement denominators.
"""
import csv
import json
from collections import Counter

from . import common

EV = common.ROOT / "evals" / "golden"
LABELS = common.ROOT / "evals" / "golden_50_labels.csv"
SEV_TOL = 0.0
SENT_TOL = 0.3  # predeclared: sentiment within ±0.3 counts as agreement


def parse_alts(s):
    out = {}
    for part in filter(None, (p.strip() for p in (s or "").split(";"))):
        if "=" in part:
            k, v = part.split("=", 1)
            out[k.strip()] = [x.strip() for x in v.split("|") if x.strip()]
    return out


LABELS_XLSX = common.ROOT / "evals" / "golden_50_labels.xlsx"


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def load_labels():
    """Read Jack's labels from the dropdown workbook if any row is filled there, else from the CSV."""
    csv_rows = list(csv.DictReader(open(LABELS, encoding="utf-8-sig", newline="")))  # CSV has a BOM for Excel
    if LABELS_XLSX.exists():
        from openpyxl import load_workbook
        ws = load_workbook(LABELS_XLSX, read_only=True)["Labels"]
        it = ws.iter_rows(values_only=True)
        hdr = [str(h) for h in next(it)]
        xrows = [{h: _cell(v) for h, v in zip(hdr, row)} for row in it if row and row[0]]
        if any(r["topic"].strip() for r in xrows):
            return xrows, LABELS_XLSX.name
    return csv_rows, LABELS.name


def load_predictions():
    import gzip
    plain, gz = common.OUT / "records.jsonl", common.ROOT / "grading" / "records.jsonl.gz"
    f = open(plain, encoding="utf-8") if plain.exists() else gzip.open(gz, "rt", encoding="utf-8")
    with f:
        return {r["review_id"]: r for r in map(json.loads, f)}


def main():
    rows, source = load_labels()
    filled = [r for r in rows if r["topic"].strip() and r["intent"].strip() and r["severity"].strip()]
    if len(filled) < len(rows):
        print(f"{len(rows) - len(filled)} of {len(rows)} golden rows are not labelled yet in {source}; "
              "fill evals/golden_50_labels.xlsx (dropdowns) or .csv first.")
        if not filled:
            return None
    print(f"labels read from {source}")
    preds = load_predictions()
    EV.mkdir(parents=True, exist_ok=True)
    out, tot = [], Counter()
    conf = {f: Counter() for f in ("topic", "intent")}
    for g in filled:
        rid, p = g["review_id"], preds.get(g["review_id"], {})
        ok = p.get("status") == "completed"
        alts = parse_alts(g.get("accepted_alternatives"))
        acc = lambda f, cast=str: [cast(g[f].strip())] + [cast(x) for x in alts.get(f, [])]
        res = {"review_id": rid, "prediction_status": p.get("status", "missing")}
        for f, cast in (("topic", str), ("intent", str), ("severity", int)):
            res[f"expected_{f}"] = "|".join(map(str, acc(f, cast)))
            res[f"predicted_{f}"] = p.get(f, "")
            res[f"{f}_pass"] = ok and p.get(f) in acc(f, cast)
            tot[f"{f}_pass"] += res[f"{f}_pass"]
        if ok:
            tot["valid"] += 1
            tot["sev_abs_err"] += abs(p["severity"] - int(g["severity"]))
            if g["sentiment"].strip():
                e = abs(p["sentiment"] - float(g["sentiment"]))
                tot["sent_n"] += 1
                tot["sent_abs_err"] += e
                res["sentiment_pass"] = e <= SENT_TOL
                tot["sentiment_pass"] += res["sentiment_pass"]
            res["quote_is_substring"] = p["evidence_quote"] in g["review_text"]
            tot["quote_substring"] += res["quote_is_substring"]
            ge = {x.strip().lower() for x in g["entities"].split(";") if x.strip()}
            res["unsupported_entities"] = "|".join(e for e in p["entities"] if e.lower() not in g["review_text"].lower())
            tot["preds_with_unsupported_entities"] += bool(res["unsupported_entities"])
            if g["needs_review"].strip():
                gt = g["needs_review"].strip().lower() == "true"
                res["needs_review_pass"] = p["needs_review"] == gt
                tot["needs_review_pass"] += res["needs_review_pass"]
                tot["needs_review_n"] += 1
            for f in conf:
                conf[f][(g[f].strip(), p[f])] += 1
        tot["ambiguous"] += g.get("ambiguous", "").strip().lower() == "true"
        res["predicted_quote"] = p.get("evidence_quote", "")
        res["quote_supports_label (hand check)"] = ""
        out.append(res)
    n = len(filled)
    summary = {"labelled_cases": n, "valid_predictions": tot["valid"],
               "topic_agreement": round(tot["topic_pass"] / n, 4), "intent_agreement": round(tot["intent_pass"] / n, 4),
               "severity_exact_agreement": round(tot["severity_pass"] / n, 4),
               "severity_mae_valid": round(tot["sev_abs_err"] / tot["valid"], 4) if tot["valid"] else None,
               "sentiment_mae_valid": round(tot["sent_abs_err"] / tot["sent_n"], 4) if tot["sent_n"] else None,
               "sentiment_within_tolerance": f"{tot['sentiment_pass']}/{tot['sent_n']} (tol ±{SENT_TOL}, predeclared)",
               "quote_exact_substring": f"{tot['quote_substring']}/{tot['valid']}",
               "predictions_with_unsupported_entities": tot["preds_with_unsupported_entities"],
               "needs_review_agreement": f"{tot['needs_review_pass']}/{tot['needs_review_n']}",
               "ambiguous_cases": tot["ambiguous"],
               "per_topic_expected": dict(Counter(g["topic"].strip() for g in filled)),
               "confusion_topic": {f"{a}->{b}": c for (a, b), c in sorted(conf["topic"].items())},
               "confusion_intent": {f"{a}->{b}": c for (a, b), c in sorted(conf["intent"].items())},
               "note": "50 cases are diagnostic, not a population estimate. Missing/quarantined predictions count as failures."}
    with open(EV / "golden_results.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(out)
    common.write_json(EV / "golden_summary.json", summary)
    print(json.dumps(summary, indent=2))
    return summary
