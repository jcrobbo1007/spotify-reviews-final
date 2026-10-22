"""Build evals/golden_50_labels.xlsx from evals/golden_50_labels.csv: same rows, with dropdowns.

  python evals/make_golden_xlsx.py                     # contract label cells blank
  python evals/make_golden_xlsx.py --draft FILE.csv    # also carry over Jack's own draft labels

With --draft, Jack's own evidence_quote, entities and needs_review are copied into the editable cells, and his
free-form topic/intent/severity/sentiment appear in locked grey "my_draft_*" columns. The contract topic, intent,
severity and sentiment cells stay blank for Jack to pick from the dropdowns. Nothing here is model-written; the
example on the Definitions tab is a made-up review, not one of the 50.
"""
import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "evals" / "golden_50_labels.csv"
OUT = ROOT / "evals" / "golden_50_labels.xlsx"
TOPICS = ["access", "usability", "playback", "downloads", "catalog", "billing", "support", "other"]
INTENTS = ["cancellation", "complaint", "request", "praise", "unclear"]
CONTRACT = ["topic", "intent", "severity", "sentiment"]
CARRIED = ["entities", "evidence_quote", "needs_review"]
OTHER_INPUTS = ["ambiguous", "accepted_alternatives", "notes"]
DRAFT = ["my_draft_topic", "my_draft_intent", "my_draft_severity", "my_draft_sentiment"]
FONT = "Arial"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # cells to fill
CARRY_FILL = PatternFill("solid", fgColor="E2EFDA")   # carried over from Jack's draft, editable
DRAFT_FILL = PatternFill("solid", fgColor="EDEDED")   # Jack's draft, locked
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--draft", type=Path, help="CSV with Jack's own draft labels (same review_ids)")
    args = ap.parse_args(argv)
    base = list(csv.DictReader(open(SRC, encoding="utf-8-sig", newline="")))
    assert all(not r[c].strip() for r in base for c in CONTRACT + CARRIED + OTHER_INPUTS), "base label columns must be blank"
    draft = {}
    if args.draft:
        draft = {r["review_id"]: r for r in csv.DictReader(open(args.draft, encoding="utf-8-sig", newline=""))
                 if r.get("review_id")}
        assert set(draft) == {r["review_id"] for r in base}, "draft must cover exactly the 50 golden IDs"
    cols = ["review_id", "review_text", "review_rating"] + (DRAFT if draft else []) + CONTRACT + CARRIED + OTHER_INPUTS
    if draft:
        cols.append("check_quote_in_text")
    wb = Workbook()
    ws = wb.active
    ws.title = "Labels"
    ws.append(cols)
    for r in base:
        d = draft.get(r["review_id"], {})
        vals = {"review_id": r["review_id"], "review_text": r["review_text"],
                "review_rating": int(r["review_rating"]) if r["review_rating"].strip().isdigit() else r["review_rating"]}
        if d:
            vals.update({"my_draft_topic": d["topic"], "my_draft_intent": d["intent"],
                         "my_draft_severity": d["severity"], "my_draft_sentiment": d["sentiment"],
                         "entities": d["entities"], "evidence_quote": d["evidence_quote"],
                         "needs_review": d["needs_review"].strip().upper()})
            q = d["evidence_quote"]
            vals["check_quote_in_text"] = ("" if not q.strip() else "OK" if q in r["review_text"]
                                           else "NOT an exact substring of review_text")
        ws.append([vals.get(c, "") for c in cols])
    n = len(base) + 1
    letter = {c: ws.cell(row=1, column=i + 1).column_letter for i, c in enumerate(cols)}
    widths = {"review_id": 38, "review_text": 64, "review_rating": 8, "topic": 13, "intent": 14, "severity": 9,
              "sentiment": 10, "entities": 24, "evidence_quote": 40, "needs_review": 13, "ambiguous": 11,
              "accepted_alternatives": 26, "notes": 28, "my_draft_topic": 26, "my_draft_intent": 16,
              "my_draft_severity": 12, "my_draft_sentiment": 13, "check_quote_in_text": 22}
    for c, w in widths.items():
        if c in letter:
            ws.column_dimensions[letter[c]].width = w
    editable = set(CONTRACT + CARRIED + OTHER_INPUTS)
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEAD_FILL
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            col = ws.cell(row=1, column=cell.column).value
            cell.font = Font(name=FONT, italic=col in DRAFT)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col in editable:
                cell.protection = Protection(locked=False)
                cell.fill = CARRY_FILL if (draft and col in CARRIED) else INPUT_FILL
            elif col in DRAFT:
                cell.fill = DRAFT_FILL
    ws.freeze_panes = "B2"

    def dv_list(col, values, prompt):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Not a valid label", error="Pick from the list.",
                            showInputMessage=True, promptTitle=col, prompt=prompt)
        dv.add(f"{letter[col]}2:{letter[col]}{n}")
        ws.add_data_validation(dv)

    dv_list("topic", TOPICS, "One primary topic (see Definitions tab).")
    dv_list("intent", INTENTS, "Precedence: cancellation > complaint > request > praise > unclear.")
    dv_list("severity", ["1", "2", "3", "4", "5"], "1 no problem ... 5 serious financial/privacy/data harm.")
    dv_list("needs_review", ["TRUE", "FALSE"], "TRUE if context is missing or the call is close.")
    dv_list("ambiguous", ["TRUE", "FALSE"], "TRUE if more than one label is defensible.")
    sent = DataValidation(type="decimal", operator="between", formula1="-1", formula2="1", allow_blank=True,
                          showErrorMessage=True, error="Sentiment must be a number between -1 and 1.",
                          showInputMessage=True, promptTitle="sentiment", prompt="-1.0 (very negative) to 1.0 (very positive).")
    sent.add(f"{letter['sentiment']}2:{letter['sentiment']}{n}")
    ws.add_data_validation(sent)
    # Lock everything except the input cells (no password; Review > Unprotect Sheet lifts it).
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.formatCells = False
    ws.protection.autoFilter = False
    ws.protection.sort = False

    d = wb.create_sheet("Definitions")
    lines = [
        ("How to use", ""),
        ("Yellow cells", "Fill these: topic, intent, severity, sentiment (dropdowns / -1..1), ambiguous, accepted_alternatives, notes."),
        ("Green cells", "Carried over from your own draft (evidence_quote, entities, needs_review). Editable; check them."),
        ("Grey 'my_draft_*' columns", "Your own draft labels, locked for reference. They are not in the contract's label set."),
        ("check_quote_in_text", "Code check only: says whether your evidence_quote is an exact substring of review_text."),
        ("Locked cells", "The sheet is protected without a password (Review > Unprotect Sheet) so review data and drafts are not edited by accident."),
        ("entities", "Feature terms the text explicitly names, separated by ';' (blank if none)."),
        ("evidence_quote", "An exact substring of review_text that supports the label."),
        ("accepted_alternatives", "Optional, e.g. topic=billing|usability; severity=2|3"),
        ("Save", "Keep this file as .xlsx; python -m pipeline eval golden reads it directly."),
        ("", ""),
        ("Example format (made-up review, NOT one of the 50)", ""),
        ("review_text", "The app logs me out every morning and I can't get back in with my password."),
        ("labels", "topic=access; intent=complaint; severity=4; sentiment=-0.7; entities=login;password; "
                   "evidence_quote=I can't get back in with my password; needs_review=FALSE; ambiguous=FALSE"),
        ("", ""),
        ("topic", "Definition (GRADING_CONTRACT.md)"),
        ("access", "Login, signup, password or account access"),
        ("usability", "Navigation, controls, layout, queue/playlist management, ad interruptions"),
        ("playback", "Playback failure, crashes, lag, connection failures, audio quality, resource use"),
        ("downloads", "Downloading, saved music, offline listening, disappearing downloads"),
        ("catalog", "Missing songs/artists, search/discovery, recommendations, lyrics availability"),
        ("billing", "Price, charges, subscriptions, paywalls, premium entitlement; explicitly premium-only controls go here"),
        ("support", "Contacting support and the support response"),
        ("other", "General praise/criticism, unrelated content, or no supported specific topic"),
        ("Topic tie rules", "Highest supported severity; on a tie, the first specific problem. Positive review: first specific "
                            "praised feature; general praise is other. A paid-plan mention alone is not billing."),
        ("", ""),
        ("intent", "cancellation (explicitly leaving, uninstalling, cancelling, or threatening to) > complaint (negative "
                   "experience, incl. mixed) > request (desired change, no reported failure) > praise > unclear. Bare "
                   "boycott slogans and meaningless text are unclear. General 'bad app' is a complaint."),
        ("", ""),
        ("severity", "Meaning"),
        ("1", "No reported problem: praise, neutral/unclear content, or a pure feature request"),
        ("2", "Dislike, generic criticism, minor annoyance, or cosmetic issue; no supported functional loss"),
        ("3", "A degraded or restricted function; some use or workaround remains"),
        ("4", "A clearly blocked core task, such as inability to log in or play music"),
        ("5", "Explicit serious financial, privacy, or data harm; an expensive plan, a crash, or angry language alone is insufficient"),
        ("Note", "Cancellation intent does not raise severity. Missing context -> needs_review; do not invent impact."),
    ]
    for a, b in lines:
        d.append([a, b])
    d.column_dimensions["A"].width = 34
    d.column_dimensions["B"].width = 110
    for row in d.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, bold=cell.column == 1 and bool(cell.value))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}: {len(base)} rows" + (f", draft carried over from {args.draft.name}" if draft else ""))


if __name__ == "__main__":
    main()
